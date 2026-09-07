#!/usr/bin/env python3
"""Normalize the official MV 2026 candidate and MV 2021 result workbooks."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_METADATA = ROOT / "data" / "2026-mv" / "metadata"
DEFAULT_REFERENCE = ROOT / "data" / "2026-mv" / "reference" / "2021"


PARTY_ALIASES = {
    "DIE LINKE": "Die Linke",
    "DIE PARTEI": "Die PARTEI",
    "FREIE  WÄHLER": "FREIE WÄHLER",
    "FREIE WÄHLER Mecklenburg-Vorpommern": "FREIE WÄHLER",
}


def text(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\xa0", " ")).strip()


def canonical_party(value: Any) -> str:
    name = text(value)
    return PARTY_ALIASES.get(name, name)


def normalized_key(value: Any) -> str:
    return " ".join(
        unicodedata.normalize("NFKD", text(value)).encode("ascii", "ignore").decode("ascii").lower().split()
    )


def integer(value: Any) -> int | None:
    if value is None or text(value).lower() in {"", "x", "-", "–"}:
        return None
    try:
        return int(float(str(value).replace(".", "").replace(",", ".")))
    except ValueError:
        return None


def percentage(votes: int | None, total: int | None) -> str:
    if votes is None or not total:
        return ""
    return f"{votes / total * 100:.6f}"


def write_csv(path: Path, fieldnames: Iterable[str], rows: Iterable[Dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(fieldnames), extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def candidate_id(party: str, surname: str, given_name: str, birth_year: Any, residence: str) -> str:
    raw = "|".join(map(normalized_key, (party, surname, given_name, birth_year, residence)))
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()[:16]


def prepare_candidates(workbook_path: Path, output_path: Path) -> Dict[str, int]:
    book = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    details: Dict[Tuple[str, str, str, str], Dict[str, str]] = {}
    wahlkreis_names: Dict[str, str] = {}

    def detail_key(party: Any, surname: Any, given_name: Any, birth_year: Any) -> Tuple[str, str, str, str]:
        return tuple(normalized_key(value) for value in (canonical_party(party), surname, given_name, birth_year))  # type: ignore[return-value]

    def remember(party: Any, surname: Any, given_name: Any, occupation: Any, birth_year: Any, residence: Any) -> None:
        key = detail_key(party, surname, given_name, birth_year)
        details.setdefault(
            key,
            {"occupation": text(occupation), "residence": text(residence)},
        )

    list_sheet = book["2. Landeslisten nach Bewerbern"]
    list_party = ""
    for values in list_sheet.iter_rows(min_row=4, values_only=True):
        first, second, given, occupation, birth_year, residence = (list(values) + [None] * 6)[:6]
        heading = text(second)
        if heading and heading.startswith("2.") and " - " in heading:
            list_party = heading.rsplit(" - ", 1)[1]
            continue
        if integer(first) is None or not list_party:
            continue
        remember(list_party, second, given, occupation, birth_year, residence)

    direct_sheet = book["4. Kreiswahlvor. n. Bewerbern"]
    current_wk = ""
    current_wk_name = ""
    for values in direct_sheet.iter_rows(min_row=4, values_only=True):
        values = list(values) + [None] * 8
        first, party, list_position, surname, given, occupation, birth_year, residence = values[:8]
        heading = text(party)
        match = re.search(r"Wahlkreis:\s*(\d+)\s*-\s*(.+)$", heading)
        if match:
            current_wk, current_wk_name = match.group(1), text(match.group(2))
            wahlkreis_names[current_wk] = current_wk_name
            continue
        if integer(first) is None or not current_wk or not text(party):
            continue
        remember(party, surname, given, occupation, birth_year, residence)

    candidates: List[Dict[str, Any]] = []
    alphabetical_sheet = book["5. Alphabetisches Verzeichnis"]
    for values in alphabetical_sheet.iter_rows(min_row=4, values_only=True):
        values = list(values) + [None] * 8
        serial, surname, given_name, party, wahlkreisnummer, list_position, birth_year, residence = values[:8]
        if integer(serial) is None:
            continue
        party = canonical_party(party)
        surname, given_name, residence = text(surname), text(given_name), text(residence)
        wahlkreisnummer = text(wahlkreisnummer)
        list_position = text(list_position)
        details_row = details.get(detail_key(party, surname, given_name, birth_year), {})
        has_wahlkreis = bool(wahlkreisnummer)
        has_list = bool(list_position)
        kind = "Landesliste+Wahlkreis" if has_list and has_wahlkreis else ("Wahlkreis" if has_wahlkreis else "Landesliste")
        candidates.append(
            {
                "candidate_id": candidate_id(party, surname, given_name, birth_year, residence),
                "party": party,
                "candidate_type": kind,
                "list_position": list_position,
                "wahlkreisnummer": wahlkreisnummer,
                "wahlkreis_name": wahlkreis_names.get(wahlkreisnummer, ""),
                "surname": surname,
                "given_name": given_name,
                "occupation": details_row.get("occupation", ""),
                "birth_year": text(birth_year),
                "residence": residence,
            }
        )
    rows = sorted(
        candidates,
        key=lambda row: (
            0 if "Wahlkreis" in row["candidate_type"] else 1,
            integer(row["wahlkreisnummer"]) or 999,
            canonical_party(row["party"]),
            integer(row["list_position"]) or 999,
            row["surname"],
            row["given_name"],
        ),
    )
    write_csv(
        output_path,
        [
            "candidate_id",
            "party",
            "candidate_type",
            "list_position",
            "wahlkreisnummer",
            "wahlkreis_name",
            "surname",
            "given_name",
            "occupation",
            "birth_year",
            "residence",
        ],
        rows,
    )
    return {
        "unique_candidates": len(rows),
        "state_list_candidates": sum("Landesliste" in row["candidate_type"] for row in rows),
        "constituency_candidates": sum("Wahlkreis" in row["candidate_type"] for row in rows),
        "constituencies": len({row["wahlkreisnummer"] for row in rows if row["wahlkreisnummer"]}),
    }


def party_codebook(book: Any) -> Dict[str, str]:
    result: Dict[str, str] = {}
    for row in book["Übersicht Parteien"].iter_rows(values_only=True):
        code, short = text(row[0]), text(row[1])
        if re.fullmatch(r"[DF]\d+", code) and short:
            result[code] = canonical_party(short)
    return result


def area_record(values: List[Any], *, municipality: bool) -> Dict[str, Any] | None:
    if municipality:
        area_id = text(values[2])
        area_name = text(values[3])
        indices = {"eligible": 7, "voters": 8, "invalid_first": 10, "valid_first": 11, "invalid_second": 45, "valid_second": 46}
    else:
        area_id = text(values[0])
        area_name = text(values[1])
        indices = {"eligible": 5, "voters": 6, "invalid_first": 8, "valid_first": 9, "invalid_second": 43, "valid_second": 44}
    if not area_id or not area_name:
        return None
    return {
        "area_id": area_id,
        "area_name": area_name,
        **{key: integer(values[index]) for key, index in indices.items()},
    }


def extract_areas(book: Any, sheet_name: str, *, municipality: bool) -> Tuple[List[Dict[str, Any]], Dict[str, Dict[str, str]]]:
    sheet = book[sheet_name]
    codebook = party_codebook(book)
    rows = list(sheet.iter_rows(min_row=12, values_only=True))
    areas: List[Dict[str, Any]] = []
    party_rows: Dict[str, Dict[str, str]] = {}
    first_party_start, first_party_end = (12, 45) if municipality else (11, 44)
    second_party_start, second_party_end = (47, 71) if municipality else (45, 69)
    for raw_values in rows:
        values = list(raw_values)
        area = area_record(values, municipality=municipality)
        if area is None:
            continue
        area["area_level"] = "GEMEINDE" if municipality else "WAHLKREIS"
        areas.append(area)
        for vote_type, start, end, prefix, valid_key in (
            ("Erststimmen", first_party_start, first_party_end, "D", "valid_first"),
            ("Zweitstimmen", second_party_start, second_party_end, "F", "valid_second"),
        ):
            for index in range(start, min(end, len(values))):
                code = f"{prefix}{index - start + 1}"
                party = codebook.get(code)
                votes = integer(values[index])
                if not party or votes is None:
                    continue
                key = f"{area['area_level']}|{area['area_id']}|{vote_type}|{party}"
                party_rows[key] = {
                    "area_level": area["area_level"],
                    "area_id": area["area_id"],
                    "area_name": area["area_name"],
                    "vote_type": vote_type,
                    "party_name": party,
                    "votes": str(votes),
                    "valid_votes": str(area[valid_key] or 0),
                }
    return areas, party_rows


def prepare_reference(workbook_path: Path, mandate_path: Path, reference_dir: Path) -> Dict[str, int]:
    book = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    wahlkreis_areas, wahlkreis_party_rows = extract_areas(book, "nach Wahlkreisen", municipality=False)
    municipality_areas, municipality_party_rows = extract_areas(book, "nach Gemeinden", municipality=True)
    all_areas = wahlkreis_areas + municipality_areas
    land = {
        "area_level": "LAND",
        "area_id": "99",
        "area_name": "Mecklenburg-Vorpommern",
        "eligible_voters": sum(area["eligible"] or 0 for area in wahlkreis_areas),
        "voters": sum(area["voters"] or 0 for area in wahlkreis_areas),
        "valid_first_votes": sum(area["valid_first"] or 0 for area in wahlkreis_areas),
        "valid_second_votes": sum(area["valid_second"] or 0 for area in wahlkreis_areas),
        "invalid_first_votes": sum(area["invalid_first"] or 0 for area in wahlkreis_areas),
        "invalid_second_votes": sum(area["invalid_second"] or 0 for area in wahlkreis_areas),
    }
    area_rows = [land]
    for area in all_areas:
        area_rows.append(
            {
                "area_level": area["area_level"],
                "area_id": area["area_id"],
                "area_name": area["area_name"],
                "eligible_voters": area["eligible"],
                "voters": area["voters"],
                "valid_first_votes": area["valid_first"],
                "valid_second_votes": area["valid_second"],
                "invalid_first_votes": area["invalid_first"],
                "invalid_second_votes": area["invalid_second"],
            }
        )
    write_csv(
        reference_dir / "areas.csv",
        ["area_level", "area_id", "area_name", "eligible_voters", "voters", "valid_first_votes", "valid_second_votes", "invalid_first_votes", "invalid_second_votes"],
        area_rows,
    )

    all_party_rows = list(wahlkreis_party_rows.values()) + list(municipality_party_rows.values())
    state_party: Dict[Tuple[str, str], int] = {}
    for row in wahlkreis_party_rows.values():
        key = (row["vote_type"], row["party_name"])
        state_party[key] = state_party.get(key, 0) + int(row["votes"])
    for (vote_type, party), votes in sorted(state_party.items()):
        valid_key = "valid_first_votes" if vote_type == "Erststimmen" else "valid_second_votes"
        all_party_rows.append(
            {
                "area_level": "LAND",
                "area_id": "99",
                "area_name": "Mecklenburg-Vorpommern",
                "vote_type": vote_type,
                "party_name": party,
                "votes": str(votes),
                "valid_votes": str(land[valid_key]),
            }
        )
    write_csv(
        reference_dir / "party_results.csv",
        ["area_level", "area_id", "area_name", "vote_type", "party_name", "votes", "valid_votes"],
        all_party_rows,
    )

    summary_rows = []
    for area in wahlkreis_areas:
        result = {
            "wahlkreisnummer": area["area_id"],
            "wahlkreis_name": area["area_name"],
            "valid_first_votes": area["valid_first"],
            "valid_second_votes": area["valid_second"],
        }
        for vote_type, field, prefix in (
            ("Erststimmen", "first", "winner_first"),
            ("Zweitstimmen", "second", "winner_second"),
        ):
            candidates = [
                row for row in wahlkreis_party_rows.values()
                if row["area_id"] == area["area_id"] and row["vote_type"] == vote_type
            ]
            winner = max(candidates, key=lambda row: int(row["votes"]), default=None)
            result[prefix] = winner["party_name"] if winner else ""
            result[f"{prefix}_votes"] = winner["votes"] if winner else ""
            result[f"{prefix}_share_percent"] = percentage(
                int(winner["votes"]) if winner else None,
                area[f"valid_{field}"] if winner else None,
            )
        summary_rows.append(result)
    write_csv(
        reference_dir / "wahlkreis_summary.csv",
        ["wahlkreisnummer", "wahlkreis_name", "winner_first", "winner_first_votes", "winner_first_share_percent", "winner_second", "winner_second_votes", "winner_second_share_percent", "valid_first_votes", "valid_second_votes"],
        summary_rows,
    )

    seats_by_party: Dict[str, Dict[str, Any]] = {}
    with mandate_path.open("r", encoding="latin-1", newline="") as handle:
        for _ in range(5):
            handle.readline()
        reader = csv.DictReader(handle, delimiter=";")
        fieldnames = reader.fieldnames or []
        party_fields = fieldnames[9:]
        for row in reader:
            mandate_type = text(row.get("Mandatstyp"))
            if text(row.get("Wahlkreis")) != "99" or mandate_type not in {"Direktmandate", "Mandate nach Landesliste", "Insgesamt"}:
                continue
            for field in party_fields:
                value = integer(row.get(field))
                if value is None:
                    continue
                party = canonical_party(field)
                seat = seats_by_party.setdefault(
                    party,
                    {"Partei": party, "Sitze gesamt": 0, "Kreiswahlvorschlaege": 0, "Landeswahlvorschlaege": 0},
                )
                if mandate_type == "Direktmandate":
                    seat["Kreiswahlvorschlaege"] = value
                elif mandate_type == "Mandate nach Landesliste":
                    seat["Landeswahlvorschlaege"] = value
                elif mandate_type == "Insgesamt":
                    seat["Sitze gesamt"] = value
    seats = [row for row in seats_by_party.values() if row["Sitze gesamt"]]
    write_csv(reference_dir / "seats.csv", ["Partei", "Sitze gesamt", "Kreiswahlvorschlaege", "Landeswahlvorschlaege"], seats)
    return {"wahlkreise": len(wahlkreis_areas), "gemeinden": len(municipality_areas), "party_rows": len(all_party_rows), "seats": len(seats)}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--candidates", type=Path, required=True)
    parser.add_argument("--results-2021", type=Path, required=True)
    parser.add_argument("--mandates-2021", type=Path, required=True)
    parser.add_argument("--metadata-dir", type=Path, default=DEFAULT_METADATA)
    parser.add_argument("--reference-dir", type=Path, default=DEFAULT_REFERENCE)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    candidate_summary = prepare_candidates(args.candidates, args.metadata_dir / "candidates.csv")
    reference_summary = prepare_reference(args.results_2021, args.mandates_2021, args.reference_dir)
    (args.metadata_dir / "candidate_summary.json").write_text(json.dumps(candidate_summary, indent=2) + "\n", encoding="utf-8")
    (args.reference_dir / "summary.json").write_text(json.dumps(reference_summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"candidates": candidate_summary, "reference_2021": reference_summary}, indent=2))


if __name__ == "__main__":
    main()
